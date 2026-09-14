import os
import json
import random
import sqlite3
import uuid
import csv
import io
from datetime import datetime, timezone
from pathlib import Path

from flask import Flask, request, jsonify, render_template, g, send_file
from openai import OpenAI
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from questions import QUESTIONS, CONDITIONS
from demographics import DEMOGRAPHIC_QUESTIONS

BASE_DIR = Path(__file__).parent
# Defaults to a file next to app.py for local dev. On a host with a persistent
# disk mounted elsewhere (e.g. Render's disk at /var/data), set DB_PATH to a
# path on that disk so the database survives restarts/redeploys.
DB_PATH = Path(os.environ.get("DB_PATH", BASE_DIR / "responses.db"))

app = Flask(__name__)

# Reads OPENAI_API_KEY from the environment. Set it before running:
#   export OPENAI_API_KEY=sk-...        (macOS/Linux)
#   $env:OPENAI_API_KEY = "sk-..."      (PowerShell)
client = OpenAI()
MODEL = os.environ.get("OPENAI_MODEL", "gpt-4o-mini")


# Qualtrics embeds this app in an iframe (see README, "Embedding inside Qualtrics").
# Without an explicit allow-list, some hosts' default headers silently block that.
# Add your Qualtrics data-center domain(s) here — check the URL of your own survey
# editor, e.g. https://yourorg.qualtrics.com or https://yourorg.az1.qualtrics.com —
# and list the exact one(s) you use.
QUALTRICS_FRAME_ANCESTORS = os.environ.get(
    "QUALTRICS_FRAME_ANCESTORS",
    "https://*.qualtrics.com",
)


@app.after_request
def allow_qualtrics_iframe(response):
    response.headers["Content-Security-Policy"] = (
        f"frame-ancestors 'self' {QUALTRICS_FRAME_ANCESTORS};"
    )
    return response


# ---------------------------------------------------------------------------
# DB helpers
# ---------------------------------------------------------------------------

def get_db():
    db = getattr(g, "_db", None)
    if db is None:
        db = g._db = sqlite3.connect(DB_PATH)
        db.row_factory = sqlite3.Row
    return db


@app.teardown_appcontext
def close_db(exc):
    db = getattr(g, "_db", None)
    if db is not None:
        db.close()


def init_db():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA journal_mode=WAL;")  # safer concurrent writes across participants
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS sessions (
            participant_id TEXT PRIMARY KEY,
            condition TEXT NOT NULL,
            started_at TEXT NOT NULL,
            finished_at TEXT
        );
        CREATE TABLE IF NOT EXISTS responses (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            participant_id TEXT NOT NULL,
            question_index INTEGER NOT NULL,
            question_text TEXT NOT NULL,
            condition TEXT NOT NULL,
            raw_answer TEXT NOT NULL,
            polished_answer TEXT NOT NULL,
            alternative_answer TEXT NOT NULL,
            final_choice TEXT NOT NULL,
            response_text TEXT NOT NULL,
            answered_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS demographics (
            participant_id TEXT PRIMARY KEY,
            answers_json TEXT NOT NULL,
            submitted_at TEXT NOT NULL
        );
        """
    )
    conn.commit()
    conn.close()


init_db()


def now_iso():
    return datetime.now(timezone.utc).isoformat()


# ---------------------------------------------------------------------------
# LLM generation
# ---------------------------------------------------------------------------

def generate_variants(question_text, raw_answer, condition_key):
    """Given the participant's own typed answer, ask the LLM for:
    1. a grammatically corrected, beautified version of THEIR SAME answer
       (same content/position, just polished wording), and
    2. a different, short alternative answer to the same statement that
       reflects the participant's (hidden) framing condition — a genuinely
       different direction from what they wrote, not just a rephrasing.
    """

    condition = CONDITIONS[condition_key]
    system_prompt = (
        "You are helping run a social-science survey about decision framing. "
        "The statement and the participant's answer are written in Greek. "
        "A participant was shown a short statement and typed their own answer. "
        "Do two things with it:\n"
        "1. Rewrite their answer as a grammatically correct, polished version, "
        "in natural, fluent Greek. Preserve their original meaning, position, "
        "and intent — fix grammar, clarity, and phrasing only. Do not add new "
        "arguments or change their stance, and do not translate it into another "
        "language.\n"
        "2. Write a second, different short answer (max ~25 words), also in "
        "natural, fluent Greek, to the same statement that reflects the "
        "following framing: "
        f"{condition['description']} "
        "This alternative must reflect a genuinely different position from the "
        "participant's own answer, not a rephrasing of it.\n"
        'Respond ONLY with JSON of the form: '
        '{"polished": "...", "alternative": "...", "rationale": "..."}. '
        "polished and alternative must be in Greek. rationale is one short "
        "sentence in English for a researcher explaining how the alternative "
        "differs from the participant's answer and why it fits the framing "
        "(it will never be shown to the participant)."
    )
    user_prompt = f"Statement: {question_text}\nParticipant's answer: {raw_answer}"

    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        temperature=0.7,
        response_format={"type": "json_object"},
    )
    raw_text = response.choices[0].message.content

    try:
        parsed = json.loads(raw_text)
        polished = parsed["polished"].strip()
        alternative = parsed["alternative"].strip()
        rationale = parsed.get("rationale", "")
        if not polished or not alternative:
            raise ValueError("empty polished/alternative")
    except Exception:
        # Never let a malformed LLM response break the survey.
        polished = raw_answer
        alternative = "Ίσως θα άξιζε μια διαφανής, ενιαία διαδικασία για όλους."
        rationale = "Fallback text used — the model response could not be parsed."

    return polished, alternative, rationale, system_prompt, user_prompt, raw_text


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------

@app.route("/")
def index():
    return render_template("index.html", total_questions=len(QUESTIONS))


@app.route("/api/session/start", methods=["POST"])
def session_start():
    data = request.get_json(force=True) or {}
    participant_id = data.get("participant_id") or str(uuid.uuid4())

    db = get_db()
    row = db.execute(
        "SELECT condition FROM sessions WHERE participant_id = ?", (participant_id,)
    ).fetchone()

    if row:
        condition = row["condition"]
    else:
        condition = random.choice(list(CONDITIONS.keys()))
        db.execute(
            "INSERT INTO sessions (participant_id, condition, started_at) VALUES (?, ?, ?)",
            (participant_id, condition, now_iso()),
        )
        db.commit()

    return jsonify(
        {
            "participant_id": participant_id,
            "condition": condition,
            "total_questions": len(QUESTIONS),
        }
    )


@app.route("/api/demographics/questions")
def get_demographic_questions():
    # Sent to the frontend so it can render the demographics form dynamically —
    # add/edit questions in demographics.py without touching any template/JS.
    return jsonify(DEMOGRAPHIC_QUESTIONS)


@app.route("/api/demographics", methods=["POST"])
def submit_demographics():
    data = request.get_json(force=True) or {}
    participant_id = data.get("participant_id")
    answers = data.get("answers")
    if not participant_id or not isinstance(answers, dict):
        return jsonify({"error": "participant_id and answers are required"}), 400

    expected_ids = {q["id"] for q in DEMOGRAPHIC_QUESTIONS}
    if not expected_ids.issubset(answers.keys()):
        return jsonify({"error": "missing answers for one or more demographic questions"}), 400

    db = get_db()
    db.execute(
        """INSERT INTO demographics (participant_id, answers_json, submitted_at)
           VALUES (?, ?, ?)
           ON CONFLICT(participant_id) DO UPDATE SET
             answers_json = excluded.answers_json,
             submitted_at = excluded.submitted_at""",
        (participant_id, json.dumps(answers, ensure_ascii=False), now_iso()),
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/question/<int:index>")
def get_question(index):
    if index < 0 or index >= len(QUESTIONS):
        return jsonify({"error": "question index out of range"}), 404
    return jsonify(
        {
            "question_index": index,
            "question_text": QUESTIONS[index],
            "total_questions": len(QUESTIONS),
        }
    )


@app.route("/api/generate_variants", methods=["POST"])
def api_generate_variants():
    data = request.get_json(force=True) or {}
    required = ["participant_id", "question_index", "question_text", "raw_answer"]
    if not all(k in data for k in required):
        return jsonify({"error": "missing fields"}), 400

    db = get_db()
    row = db.execute(
        "SELECT condition FROM sessions WHERE participant_id = ?",
        (data["participant_id"],),
    ).fetchone()
    if not row:
        return jsonify({"error": "unknown participant_id — call /api/session/start first"}), 400
    condition_key = row["condition"]

    polished, alternative, rationale, sys_prompt, user_prompt, raw_text = generate_variants(
        data["question_text"], data["raw_answer"], condition_key
    )

    return jsonify(
        {
            "polished": polished,
            "alternative": alternative,
            "generation_debug": {
                "model": MODEL,
                "system_prompt": sys_prompt,
                "user_prompt": user_prompt,
                "raw_response": raw_text,
                "rationale": rationale,
            },
        }
    )


@app.route("/api/answer", methods=["POST"])
def submit_answer():
    data = request.get_json(force=True) or {}
    required = [
        "participant_id",
        "question_index",
        "question_text",
        "raw_answer",
        "polished_answer",
        "alternative_answer",
        "final_choice",
        "response_text",
    ]
    if not all(k in data for k in required):
        return jsonify({"error": "missing fields"}), 400
    if data["final_choice"] not in ("own", "polished", "alternative"):
        return jsonify({"error": "final_choice must be own, polished, or alternative"}), 400

    db = get_db()
    row = db.execute(
        "SELECT condition FROM sessions WHERE participant_id = ?", (data["participant_id"],)
    ).fetchone()
    condition_key = row["condition"] if row else "unknown"

    db.execute(
        """INSERT INTO responses
           (participant_id, question_index, question_text, condition, raw_answer,
            polished_answer, alternative_answer, final_choice, response_text, answered_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            data["participant_id"],
            data["question_index"],
            data["question_text"],
            condition_key,
            data["raw_answer"],
            data["polished_answer"],
            data["alternative_answer"],
            data["final_choice"],
            data["response_text"],
            now_iso(),
        ),
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/session/finish", methods=["POST"])
def session_finish():
    data = request.get_json(force=True) or {}
    participant_id = data.get("participant_id")
    if not participant_id:
        return jsonify({"error": "participant_id is required"}), 400

    db = get_db()
    db.execute(
        "UPDATE sessions SET finished_at = ? WHERE participant_id = ?",
        (now_iso(), participant_id),
    )
    db.commit()

    rows = db.execute(
        "SELECT final_choice FROM responses WHERE participant_id = ?",
        (participant_id,),
    ).fetchall()
    total = len(rows)
    chose_alternative = sum(1 for r in rows if r["final_choice"] == "alternative")
    condition_row = db.execute(
        "SELECT condition FROM sessions WHERE participant_id = ?", (participant_id,)
    ).fetchone()

    return jsonify(
        {
            "participant_id": participant_id,
            "condition": condition_row["condition"] if condition_row else None,
            "answered": total,
            "chose_alternative": chose_alternative,
            "alternative_rate": (chose_alternative / total) if total else None,
        }
    )


def fetch_demographics_map(db):
    """Returns {participant_id: {field_id: answer, ...}} for every participant
    who has submitted the demographics form — used to denormalize demographic
    columns directly onto each response row in the exports."""
    rows = db.execute("SELECT participant_id, answers_json FROM demographics").fetchall()
    return {r["participant_id"]: json.loads(r["answers_json"]) for r in rows}


@app.route("/api/export.csv")
def export_csv():
    db = get_db()
    rows = db.execute(
        "SELECT * FROM responses ORDER BY participant_id, question_index"
    ).fetchall()
    demo_field_ids = [q["id"] for q in DEMOGRAPHIC_QUESTIONS]
    demo_map = fetch_demographics_map(db)

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(
        [
            "participant_id",
            "question_index",
            "question_text",
            "condition",
            "raw_answer",
            "polished_answer",
            "alternative_answer",
            "final_choice",
            "response_text",
            "answered_at",
        ]
        + demo_field_ids
    )
    for r in rows:
        demo_answers = demo_map.get(r["participant_id"], {})
        writer.writerow(
            [
                r["participant_id"],
                r["question_index"],
                r["question_text"],
                r["condition"],
                r["raw_answer"],
                r["polished_answer"],
                r["alternative_answer"],
                r["final_choice"],
                r["response_text"],
                r["answered_at"],
            ]
            + [demo_answers.get(fid, "") for fid in demo_field_ids]
        )
    return buf.getvalue(), 200, {"Content-Type": "text/csv"}


@app.route("/api/export_demographics.csv")
def export_demographics_csv():
    db = get_db()
    rows = db.execute(
        "SELECT * FROM demographics ORDER BY participant_id"
    ).fetchall()
    field_ids = [q["id"] for q in DEMOGRAPHIC_QUESTIONS]

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["participant_id"] + field_ids + ["submitted_at"])
    for r in rows:
        answers = json.loads(r["answers_json"])
        writer.writerow(
            [r["participant_id"]] + [answers.get(fid, "") for fid in field_ids] + [r["submitted_at"]]
        )
    return buf.getvalue(), 200, {"Content-Type": "text/csv"}


@app.route("/api/export.xlsx")
def export_xlsx():
    db = get_db()
    rows = db.execute(
        "SELECT * FROM responses ORDER BY participant_id, question_index"
    ).fetchall()

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="204D53", end_color="204D53", fill_type="solid")
    body_font = Font(name="Arial")
    wrap = Alignment(vertical="top", wrap_text=True)

    demo_field_ids = [q["id"] for q in DEMOGRAPHIC_QUESTIONS]
    demo_map = fetch_demographics_map(db)

    wb = Workbook()

    # ---- sheet 1: raw responses (demographic columns denormalized onto each row,
    # so no join/VLOOKUP is needed to analyze responses alongside demographics) ----
    ws = wb.active
    ws.title = "responses"
    headers = [
        "participant_id",
        "question_index",
        "question_text",
        "bias_condition",       # the randomly assigned framing — the bias metric
        "raw_answer",           # what the participant typed
        "polished_answer",      # AI grammar-corrected/beautified version of raw_answer
        "alternative_answer",   # AI-generated answer biased in the assigned direction
        "final_choice",         # own | polished | alternative
        "response_text",        # the text actually recorded as their final answer
        "answered_at",
    ] + demo_field_ids
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for r in rows:
        demo_answers = demo_map.get(r["participant_id"], {})
        ws.append(
            [
                r["participant_id"],
                r["question_index"],
                r["question_text"],
                r["condition"],
                r["raw_answer"],
                r["polished_answer"],
                r["alternative_answer"],
                r["final_choice"],
                r["response_text"],
                r["answered_at"],
            ]
            + [demo_answers.get(fid, "") for fid in demo_field_ids]
        )

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = wrap

    widths = [24, 14, 40, 14, 34, 34, 34, 12, 34, 24] + [20] * len(demo_field_ids)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    # ---- sheet 2: summary by condition ----
    ws2 = wb.create_sheet("summary_by_condition")
    summary_headers = [
        "bias_condition",
        "responses",
        "kept_own",
        "chose_polished",
        "chose_alternative",
        "alternative_rate",
    ]
    ws2.append(summary_headers)
    for col in range(1, len(summary_headers) + 1):
        cell = ws2.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill

    agg = db.execute(
        """SELECT condition,
                  COUNT(*) AS n,
                  SUM(CASE WHEN final_choice = 'own' THEN 1 ELSE 0 END) AS kept_own,
                  SUM(CASE WHEN final_choice = 'polished' THEN 1 ELSE 0 END) AS chose_polished,
                  SUM(CASE WHEN final_choice = 'alternative' THEN 1 ELSE 0 END) AS chose_alternative
           FROM responses
           GROUP BY condition"""
    ).fetchall()
    for a in agg:
        alt_rate = (a["chose_alternative"] / a["n"]) if a["n"] else 0
        ws2.append(
            [a["condition"], a["n"], a["kept_own"], a["chose_polished"], a["chose_alternative"], alt_rate]
        )

    for row in ws2.iter_rows(min_row=2, min_col=6, max_col=6):
        for cell in row:
            cell.number_format = "0.0%"
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
    for i, w in enumerate([16, 12, 12, 14, 16, 14], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ---- sheet 3: demographics ----
    ws3 = wb.create_sheet("demographics")
    field_ids = [q["id"] for q in DEMOGRAPHIC_QUESTIONS]
    demo_headers = ["participant_id"] + field_ids + ["submitted_at"]
    ws3.append(demo_headers)
    for col in range(1, len(demo_headers) + 1):
        cell = ws3.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill

    demo_rows = db.execute(
        "SELECT * FROM demographics ORDER BY participant_id"
    ).fetchall()
    for r in demo_rows:
        answers = json.loads(r["answers_json"])
        ws3.append(
            [r["participant_id"]] + [answers.get(fid, "") for fid in field_ids] + [r["submitted_at"]]
        )

    for row in ws3.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
    demo_widths = [24] + [22] * len(field_ids) + [24]
    for i, w in enumerate(demo_widths, start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="bias_survey_responses.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    app.run(debug=True, port=5000)
